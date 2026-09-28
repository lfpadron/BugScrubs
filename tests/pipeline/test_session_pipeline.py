from __future__ import annotations

from pathlib import Path
import shutil
from uuid import uuid4

from openpyxl import Workbook

from bugscrub.bug_engine.service import BugEngine
from bugscrub.db.duckdb_store import DuckDBStore
from bugscrub.discrepancies.service import compare_inventory_vs_parsed
from bugscrub.normalization.service import Normalizer, generate_inventory_workbook_from_parsed_records
from bugscrub.parsers.service import parse_runtime_session


FIXTURE_ROOT = Path(__file__).resolve().parents[1] / "fixtures"


def test_nexus_session_is_normalized_and_persisted() -> None:
    workspace_dir = build_workspace_dir("nexus")
    session_dir = build_runtime_session(workspace_dir, fixture_family="nexus")

    try:
        parsed_session = parse_runtime_session("Nexus", session_dir)
        batch = Normalizer().normalize_parsed_device(
            parsed_record=parsed_session.parsed_record,
            platform_family=parsed_session.platform_family,
            support_level=parsed_session.support_level,
            session_dir=session_dir,
            inventory_path=session_dir / "inventory.xlsx",
            source_count=5,
            warnings=parsed_session.warnings,
        )
        batch.discrepancy_rows = compare_inventory_vs_parsed(
            session_id=batch.session_id,
            inventory_rows=batch.inventory_rows,
            parsed_devices=batch.devices,
        )
        batch.bug_findings = [finding.to_record() for finding in BugEngine().run(batch)]

        assert batch.platform_family == "nexus"
        assert batch.support_level == "robust"
        assert batch.inventory_workbook.file_name == "inventory.xlsx"
        assert batch.inventory_workbook.sheet_count == 2
        assert batch.inventory_workbook.selected_sheet == "Inventory_Input"
        assert "hostname" in batch.inventory_workbook.normalized_columns
        assert batch.inventory_workbook.parsed_row_count == 2
        assert len(batch.inventory_rows) == 2
        assert batch.inventory_rows[0].hostname == "leaf01"
        assert batch.inventory_rows[0].pid == "N9K-C93180YC-FX"
        assert batch.inventory_rows[0].os_version == "9.3(9)"
        assert batch.inventory_rows[0].target_version == "10.2(5)"
        assert "VXLAN" in batch.inventory_rows[0].features
        assert len(batch.discrepancy_rows) == 4
        assert batch.discrepancy_rows[0].row_source == "cliente"
        assert batch.discrepancy_rows[1].row_source == "descubierto"
        assert batch.discrepancy_rows[0].discrepancy_status == "discrepancia"
        assert batch.discrepancy_rows[2].discrepancy_status == "faltante"
        assert len(batch.bug_findings) >= 1
        assert batch.bug_findings[0]["bug_id"] == "CSCvx10001"
        assert batch.devices[0].hostname == "leaf01"
        assert batch.devices[0].os_name == "NX-OS"
        assert batch.devices[0].os_version == "9.3(9)"

        store = DuckDBStore(workspace_dir / "nexus.duckdb")
        store.save_inventory_batch(batch)

        sessions = store.fetch_sessions()
        devices = store.fetch_devices_for_session(batch.session_id)

        assert len(sessions) == 1
        assert sessions[0]["platform_family"] == "nexus"
        assert sessions[0]["source_count"] == 5
        assert len(devices) == 1
        assert devices[0]["hostname"] == "leaf01"
        assert "VXLAN" in devices[0]["features"]
        inventory_rows = store.fetch_inventory_rows_for_session(batch.session_id)
        assert len(inventory_rows) == 2
        assert inventory_rows[1]["hostname"] == "leaf02"
        assert inventory_rows[1]["role"] == "spine"
        discrepancy_rows = store.fetch_discrepancy_rows_for_session(batch.session_id)
        assert len(discrepancy_rows) == 4
        assert discrepancy_rows[1]["discrepancy_status"] == "discrepancia"
        assert discrepancy_rows[3]["discrepancy_status"] == "faltante"
        bug_findings = store.fetch_bug_findings_for_session(batch.session_id)
        assert len(bug_findings) >= 1
        assert bug_findings[0]["hostname"] == "leaf01"
    finally:
        shutil.rmtree(workspace_dir, ignore_errors=True)


def test_catalyst_session_is_normalized_and_persisted() -> None:
    workspace_dir = build_workspace_dir("catalyst")
    session_dir = build_runtime_session(workspace_dir, fixture_family="catalyst")

    try:
        parsed_session = parse_runtime_session("Catalyst", session_dir)
        batch = Normalizer().normalize_parsed_device(
            parsed_record=parsed_session.parsed_record,
            platform_family=parsed_session.platform_family,
            support_level=parsed_session.support_level,
            session_dir=session_dir,
            inventory_path=session_dir / "inventory.xlsx",
            source_count=5,
            warnings=parsed_session.warnings,
        )
        batch.discrepancy_rows = compare_inventory_vs_parsed(
            session_id=batch.session_id,
            inventory_rows=batch.inventory_rows,
            parsed_devices=batch.devices,
        )
        batch.bug_findings = [finding.to_record() for finding in BugEngine().run(batch)]

        assert batch.platform_family == "catalyst"
        assert batch.support_level == "basic"
        assert batch.inventory_workbook.file_name == "inventory.xlsx"
        assert batch.inventory_workbook.selected_sheet == "Devices"
        assert batch.inventory_workbook.parsed_row_count == 2
        assert len(batch.inventory_rows) == 2
        assert batch.inventory_rows[0].hostname == "dist-sw01"
        assert batch.inventory_rows[0].model == "C9300-48P"
        assert batch.inventory_rows[0].os_version == "17.9.4a"
        assert batch.inventory_rows[0].site == "HQ"
        assert "EtherChannel" in batch.inventory_rows[0].features
        assert len(batch.discrepancy_rows) == 4
        assert batch.discrepancy_rows[0].discrepancy_status == "discrepancia"
        assert batch.discrepancy_rows[2].discrepancy_status == "faltante"
        assert len(batch.bug_findings) >= 1
        assert batch.bug_findings[0]["platform_family"] == "catalyst"
        assert batch.devices[0].hostname == "dist-sw01"
        assert batch.devices[0].os_name == "IOS XE"
        assert batch.devices[0].os_version == "17.9.4a"

        store = DuckDBStore(workspace_dir / "catalyst.duckdb")
        store.save_inventory_batch(batch)

        sessions = store.fetch_sessions()
        devices = store.fetch_devices_for_session(batch.session_id)

        assert len(sessions) == 1
        assert sessions[0]["platform_family"] == "catalyst"
        assert len(devices) == 1
        assert devices[0]["hostname"] == "dist-sw01"
        assert "EtherChannel" in devices[0]["features"]
        inventory_rows = store.fetch_inventory_rows_for_session(batch.session_id)
        assert len(inventory_rows) == 2
        assert inventory_rows[1]["site"] == "Branch"
        assert inventory_rows[1]["business_criticality"] == 2
        discrepancy_rows = store.fetch_discrepancy_rows_for_session(batch.session_id)
        assert len(discrepancy_rows) == 4
        assert discrepancy_rows[0]["discrepancy_status"] == "discrepancia"
        assert discrepancy_rows[2]["discrepancy_status"] == "faltante"
        bug_findings = store.fetch_bug_findings_for_session(batch.session_id)
        assert len(bug_findings) >= 1
        assert bug_findings[0]["platform_family"] == "catalyst"
    finally:
        shutil.rmtree(workspace_dir, ignore_errors=True)


def test_nexus_multi_device_session_is_normalized_and_persisted() -> None:
    workspace_dir = build_workspace_dir("nexus-multi")
    session_dir = build_multi_device_runtime_session(workspace_dir)

    try:
        parsed_session = parse_runtime_session("Nexus", session_dir)
        batch = Normalizer().normalize_parsed_devices(
            parsed_records=parsed_session.parsed_records,
            platform_family=parsed_session.platform_family,
            support_level=parsed_session.support_level,
            session_dir=session_dir,
            inventory_path=session_dir / "inventory.xlsx",
            source_count=2,
            warnings=parsed_session.warnings,
        )
        batch.discrepancy_rows = compare_inventory_vs_parsed(
            session_id=batch.session_id,
            inventory_rows=batch.inventory_rows,
            parsed_devices=batch.devices,
        )
        batch.bug_findings = [finding.to_record() for finding in BugEngine().run(batch)]

        assert batch.platform_family == "nexus"
        assert batch.support_level == "robust"
        assert len(parsed_session.parsed_records) == 2
        assert sorted(record["bundle_name"] for record in parsed_session.parsed_records) == ["leaf01", "leaf02"]
        assert len(batch.devices) == 2
        assert sorted(device.hostname for device in batch.devices) == ["leaf01", "leaf02"]
        assert len(batch.inventory_rows) == 2
        assert len(batch.discrepancy_rows) == 2
        assert {row.discrepancy_status for row in batch.discrepancy_rows} == {"discrepancia"}
        assert {row.row_source for row in batch.discrepancy_rows} == {"cliente", "descubierto"}
        assert len(batch.bug_findings) >= 2
        assert {"leaf01", "leaf02"}.issubset({finding["hostname"] for finding in batch.bug_findings})

        store = DuckDBStore(workspace_dir / "nexus-multi.duckdb")
        store.save_inventory_batch(batch)

        sessions = store.fetch_sessions()
        devices = store.fetch_devices_for_session(batch.session_id)
        discrepancy_rows = store.fetch_discrepancy_rows_for_session(batch.session_id)
        bug_findings = store.fetch_bug_findings_for_session(batch.session_id)

        assert len(sessions) == 1
        assert sessions[0]["source_count"] == 2
        assert len(devices) == 2
        assert sorted(device["hostname"] for device in devices) == ["leaf01", "leaf02"]
        assert len(discrepancy_rows) == 2
        assert all(row["discrepancy_status"] == "discrepancia" for row in discrepancy_rows)
        assert len(bug_findings) >= 2
        assert {"leaf01", "leaf02"}.issubset({finding["hostname"] for finding in bug_findings})
    finally:
        shutil.rmtree(workspace_dir, ignore_errors=True)


def test_nexus_session_without_inventory_generates_inventory_from_show_outputs() -> None:
    workspace_dir = build_workspace_dir("nexus-generated")
    session_dir = build_runtime_session(workspace_dir, fixture_family="nexus", create_inventory=False)

    try:
        parsed_session = parse_runtime_session("Nexus", session_dir)
        generated_inventory_path = generate_inventory_workbook_from_parsed_records(
            parsed_records=parsed_session.parsed_records,
            destination=session_dir / "generated_inventory.xlsx",
            platform_family=parsed_session.platform_family,
        )
        batch = Normalizer().normalize_parsed_devices(
            parsed_records=parsed_session.parsed_records,
            platform_family=parsed_session.platform_family,
            support_level=parsed_session.support_level,
            session_dir=session_dir,
            inventory_path=generated_inventory_path,
            source_count=4,
            warnings=parsed_session.warnings,
        )
        batch.discrepancy_rows = compare_inventory_vs_parsed(
            session_id=batch.session_id,
            inventory_rows=batch.inventory_rows,
            parsed_devices=batch.devices,
        )
        batch.bug_findings = [finding.to_record() for finding in BugEngine().run(batch)]

        assert generated_inventory_path.exists()
        assert batch.inventory_workbook.file_name == "generated_inventory.xlsx"
        assert batch.inventory_workbook.selected_sheet == "Inventory_Input"
        assert batch.inventory_workbook.parsed_row_count == 1
        assert len(batch.inventory_rows) == 1
        assert batch.inventory_rows[0].hostname == "leaf01"
        assert batch.inventory_rows[0].model == "Nexus9000 C93180YC-FX"
        assert batch.inventory_rows[0].pid == "N9K-C93180YC-FX"
        assert batch.inventory_rows[0].serial == "FDO1234ABCD"
        assert batch.inventory_rows[0].family == "nexus9000"
        assert batch.inventory_rows[0].features == ["VXLAN", "EVPN", "BGP", "VPC", "OSPF", "PBR", "QoS"]
        assert batch.discrepancy_rows == []
        assert len(batch.bug_findings) >= 1

        store = DuckDBStore(workspace_dir / "nexus-generated.duckdb")
        store.save_inventory_batch(batch)

        inventory_rows = store.fetch_inventory_rows_for_session(batch.session_id)
        devices = store.fetch_devices_for_session(batch.session_id)

        assert len(inventory_rows) == 1
        assert inventory_rows[0]["hostname"] == "leaf01"
        assert len(devices) == 1
        assert devices[0]["inventory_file"] == "generated_inventory.xlsx"
    finally:
        shutil.rmtree(workspace_dir, ignore_errors=True)


def build_runtime_session(tmp_path: Path, *, fixture_family: str, create_inventory: bool = True) -> Path:
    fixture_dir = FIXTURE_ROOT / fixture_family
    session_dir = tmp_path / "runtime" / fixture_family / "session-001"
    session_dir.mkdir(parents=True, exist_ok=True)

    shutil.copy2(fixture_dir / "show_version.txt", session_dir / "show_version.txt")
    shutil.copy2(fixture_dir / "show_inventory.txt", session_dir / "show_inventory.txt")
    shutil.copy2(fixture_dir / "show_running_config.txt", session_dir / "show_running_config.txt")
    shutil.copy2(fixture_dir / "show_module.txt", session_dir / "show_module.txt")
    if create_inventory:
        create_inventory_workbook(session_dir / "inventory.xlsx", fixture_family=fixture_family)
    return session_dir


def build_multi_device_runtime_session(tmp_path: Path) -> Path:
    fixture_dir = FIXTURE_ROOT / "nexus"
    session_dir = tmp_path / "runtime" / "nexus-multi" / "session-001"
    leaf01_dir = session_dir / "bundles" / "leaf01"
    leaf02_dir = session_dir / "bundles" / "leaf02"
    leaf01_dir.mkdir(parents=True, exist_ok=True)
    leaf02_dir.mkdir(parents=True, exist_ok=True)

    copy_nexus_fixture_bundle(fixture_dir=fixture_dir, destination=leaf01_dir)
    write_leaf02_nexus_bundle(fixture_dir=fixture_dir, destination=leaf02_dir)
    create_multi_device_inventory_workbook(session_dir / "inventory.xlsx")
    return session_dir


def build_workspace_dir(prefix: str) -> Path:
    workspace_dir = Path("storage") / "test-artifacts" / f"{prefix}-{uuid4().hex[:8]}"
    workspace_dir.mkdir(parents=True, exist_ok=True)
    return workspace_dir


def create_inventory_workbook(destination: Path, *, fixture_family: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if fixture_family == "nexus":
        sheet.title = "Inventory_Input"
        sheet.append(
            [
                "Device Name",
                "Platform PID",
                "Current Version",
                "Target Version",
                "Role",
                "Site",
                "Family",
                "Features",
                "Business Criticality",
            ]
        )
        sheet.append(["leaf01", "N9K-C93180YC-FX", "9.3(9)", "10.2(5)", "leaf", "DC1", "nexus9000", "VXLAN; EVPN", 1])
        sheet.append(["leaf02", "N9K-C93180YC-FX", "9.3(8)", "10.2(5)", "spine", "DC1", "nexus9000", "BGP / OSPF", 3])
    else:
        sheet.title = "Devices"
        sheet.append(
            [
                "Host Name",
                "Model Number",
                "Serial Number",
                "OS Version",
                "Location",
                "Device Role",
                "Product Family",
                "Enabled Features",
                "Criticality",
            ]
        )
        sheet.append(["dist-sw01", "C9300-48P", "FCW2145L0AB", "17.9.4a", "HQ", "distribution", "catalyst9300", "STP, EtherChannel", 1])
        sheet.append(["dist-sw02", "C9300-48P", "FCW2145L0AC", "17.9.4a", "Branch", "distribution", "catalyst9300", "VLAN; QoS", 2])
    workbook.create_sheet("Candidate_Bugs")
    workbook.save(destination)


def create_multi_device_inventory_workbook(destination: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory_Input"
    sheet.append(
        [
            "Device Name",
            "Platform PID",
            "Serial Number",
            "Current Version",
            "Target Version",
            "Role",
            "Site",
            "Family",
            "Features",
            "Business Criticality",
        ]
    )
    sheet.append(
        [
            "leaf01",
            "N9K-C93180YC-FX",
            "FDO1234ABCD",
            "9.3(9)",
            "10.2(5)",
            "leaf",
            "DC1",
            "nexus9000",
            "VXLAN; EVPN",
            1,
        ]
    )
    sheet.append(
        [
            "leaf02",
            "N9K-C93180YC-FX",
            "FDO5678EFGH",
            "9.3(8)",
            "10.2(5)",
            "spine",
            "DC1",
            "nexus9000",
            "BGP / OSPF",
            3,
        ]
    )
    workbook.create_sheet("Candidate_Bugs")
    workbook.save(destination)


def copy_nexus_fixture_bundle(*, fixture_dir: Path, destination: Path) -> None:
    shutil.copy2(fixture_dir / "show_version.txt", destination / "show_version.txt")
    shutil.copy2(fixture_dir / "show_inventory.txt", destination / "show_inventory.txt")
    shutil.copy2(fixture_dir / "show_running_config.txt", destination / "show_running_config.txt")
    shutil.copy2(fixture_dir / "show_module.txt", destination / "show_module.txt")


def write_leaf02_nexus_bundle(*, fixture_dir: Path, destination: Path) -> None:
    replacements = {
        "leaf01": "leaf02",
        "FDO1234ABCD": "FDO5678EFGH",
        "9.3(9)": "9.3(8)",
    }
    write_fixture_with_replacements(
        source=fixture_dir / "show_version.txt",
        destination=destination / "show_version.txt",
        replacements=replacements,
    )
    write_fixture_with_replacements(
        source=fixture_dir / "show_inventory.txt",
        destination=destination / "show_inventory.txt",
        replacements=replacements,
    )
    write_fixture_with_replacements(
        source=fixture_dir / "show_module.txt",
        destination=destination / "show_module.txt",
        replacements=replacements,
    )
    (destination / "show_running_config.txt").write_text(
        "\n".join(
            [
                "hostname leaf02",
                "feature ospf",
                "feature bgp",
                "",
                "router ospf UNDERLAY",
                "router bgp 65001",
                "",
            ]
        ),
        encoding="utf-8",
    )


def write_fixture_with_replacements(*, source: Path, destination: Path, replacements: dict[str, str]) -> None:
    content = source.read_text(encoding="utf-8")
    for original, replacement in replacements.items():
        content = content.replace(original, replacement)
    destination.write_text(content, encoding="utf-8")
