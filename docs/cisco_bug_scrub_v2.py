#!/usr/bin/env python3
"""
Cisco bug scrub helper v2.

What is new in v2
-----------------
- Inventory filtering by site, role, or family.
- Per-device ranking sheet.
- Go / No-Go upgrade recommendation per node.
- Optional CSV exports for scored rows and device summary.
- Backward compatible with the original workbook format.

Workbook mode
-------------
Reads these sheets:
- Inventory_Input
- Candidate_Bugs

Writes / refreshes these sheets:
- Auto_Scored
- Exec_Summary_Auto
- Device_Ranking
- Upgrade_Decision

Optional API pull mode
----------------------
Pulls bugs from the Cisco Bug API into CSV if you have valid credentials.

Examples
--------
Workbook mode:
python cisco_bug_scrub_v2.py \
  --workbook cisco_bug_scrub_nexus6_simulation.xlsx \
  --output cisco_bug_scrub_nexus6_simulation_scored_v2.xlsx

Workbook mode with filters:
python cisco_bug_scrub_v2.py \
  --workbook cisco_bug_scrub_nexus6_simulation.xlsx \
  --output filtered.xlsx \
  --site DC1 \
  --role spine,leaf \
  --family nexus9000

Workbook mode with CSV exports:
python cisco_bug_scrub_v2.py \
  --workbook cisco_bug_scrub_nexus6_simulation.xlsx \
  --output out.xlsx \
  --export-scored-csv auto_scored.csv \
  --export-device-csv device_summary.csv

API pull mode:
export CISCO_CLIENT_ID=...
export CISCO_CLIENT_SECRET=...
python cisco_bug_scrub_v2.py \
  --api-product-series "Cisco Nexus 9000 Series Switches" \
  --api-affected-release "10.4(4)M" \
  --api-output nexus_1044_bugs.csv
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import quote

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


TOKEN_URL = "https://id.cisco.com/oauth2/default/v1/token"
BUG_API_BASE = "https://apix.cisco.com/bug/v2.0/bugs"
HIGH_FILL = "F4CCCC"
MED_FILL = "FFF2CC"
LOW_FILL = "D9EAD3"
HEADER_FILL = "1F4E78"
SUBHEADER_FILL = "D9E2F3"


@dataclass
class Device:
    device_id: str
    hostname: str
    platform_pid: str
    base_pid: str
    current_version: str
    target_version: str
    role: str
    features: List[str]
    business_criticality: int
    site: str = ""
    family: str = ""


@dataclass
class Bug:
    bug_id: str
    headline: str
    source_type: str
    product_scope: str
    affected_releases: str
    fixed_releases: str
    status: str
    severity_public: str
    trigger_features: List[str]
    trigger_platforms: List[str]
    impact_type: str
    recommended_action: str
    source_url: str


@dataclass
class DeviceSummary:
    device_id: str
    hostname: str
    site: str
    family: str
    role: str
    platform_pid: str
    current_version: str
    target_version: str
    business_criticality: int
    findings_total: int
    high_count: int
    medium_count: int
    low_count: int
    max_score: int
    weighted_risk_score: int
    top_bug_id: str
    top_bug_headline: str
    top_action: str
    upgrade_decision: str
    decision_reason: str


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


def split_csvish(value: str) -> List[str]:
    if not value:
        return []
    parts = re.split(r"[;,/]|(?:\s+\+\s+)", value)
    return [norm(p) for p in parts if norm(p)]


def parse_int(value: str, default: int = 3) -> int:
    try:
        return int(str(value).strip())
    except Exception:
        return default


def extract_version_tokens(text: str) -> List[str]:
    return re.findall(r"\d+\.\d+\(\d+\)[A-Za-z0-9\-]*", text or "")


def version_key(version: str) -> Tuple[int, int, int, str]:
    m = re.match(r"(\d+)\.(\d+)\((\d+)\)([A-Za-z0-9\-]*)", version.strip())
    if not m:
        return (0, 0, 0, "")
    return (int(m.group(1)), int(m.group(2)), int(m.group(3)), m.group(4) or "")


def version_in_text(current_version: str, text: str) -> bool:
    current_version = current_version.strip()
    if not current_version or not text:
        return False
    if current_version in text:
        return True

    ranges = re.findall(
        r"(\d+\.\d+\(\d+\)[A-Za-z0-9\-]*)\s+through\s+(\d+\.\d+\(\d+\)[A-Za-z0-9\-]*)",
        text,
    )
    if ranges:
        cur = version_key(current_version)
        for start, end in ranges:
            if version_key(start) <= cur <= version_key(end):
                return True
    return False


def derive_family(platform_pid: str, base_pid: str, product_scope: str = "") -> str:
    hay = " ".join([platform_pid or "", base_pid or "", product_scope or ""]).lower()
    if "n9k" in hay or "nexus 9000" in hay:
        return "nexus9000"
    if "n7k" in hay or "nexus 7000" in hay:
        return "nexus7000"
    if "n5k" in hay or "nexus 5000" in hay:
        return "nexus5000"
    if "n3k" in hay or "nexus 3000" in hay:
        return "nexus3000"
    if "nexus" in hay:
        return "nexus"
    return "unknown"


def platform_match(device_pid: str, bug: Bug) -> int:
    if not bug.trigger_platforms:
        return 1
    pid = norm(device_pid)
    for platform in bug.trigger_platforms:
        if platform and platform in pid:
            return 3
    scope = norm(bug.product_scope)
    if "nexus 9000" in scope and "n9k" in pid:
        return 1
    if "nexus 7000" in scope and "n7k" in pid:
        return 1
    if "nexus 5000" in scope and "n5k" in pid:
        return 1
    if "nexus 3000" in scope and "n3k" in pid:
        return 1
    return 0


def feature_match(device_features: Sequence[str], bug: Bug) -> int:
    if not bug.trigger_features:
        return 0
    device_set = set(device_features)
    count = 0
    for trig in bug.trigger_features:
        if not trig:
            continue
        if trig in device_set:
            count += 1
            continue
        for feat in device_set:
            if trig in feat or feat in trig:
                count += 1
                break
    return count


def severity_score(severity_text: str) -> int:
    m = re.search(r"([1-6])", severity_text or "")
    if not m:
        return 0
    sev = int(m.group(1))
    return max(0, 6 - sev)


def severity_numeric(severity_text: str) -> int:
    m = re.search(r"([1-6])", severity_text or "")
    if not m:
        return 6
    return int(m.group(1))


def status_score(status_text: str) -> int:
    st = norm(status_text)
    if "open" in st:
        return 2
    if "fixed" in st or "resolved" in st:
        return 0
    return 1


def band_from_score(score: int) -> str:
    if score >= 9:
        return "High"
    if score >= 6:
        return "Medium"
    return "Low"


def action_hint(device: Device, bug: Bug, score: int, version_match: bool) -> str:
    rec = norm(bug.recommended_action)
    if version_match and ("upgrade" in rec or bug.fixed_releases.strip()):
        return bug.recommended_action or f"Review fixed releases: {bug.fixed_releases}"
    if score >= 9:
        return "Review immediately; likely action required before or during upgrade."
    if score >= 6:
        return "Validate trigger conditions and decide whether to mitigate or move target release."
    return "Monitor / document unless the feature is business-critical."


def get_header_map(ws, header_row: int = 2) -> Dict[str, int]:
    return {str(cell.value).strip(): idx + 1 for idx, cell in enumerate(ws[header_row]) if cell.value}


def optional_cell(ws, row: int, header: Dict[str, int], column_name: str) -> str:
    idx = header.get(column_name)
    if not idx:
        return ""
    return str(ws.cell(row, idx).value or "").strip()


def load_devices(ws) -> List[Device]:
    header = get_header_map(ws)
    required = [
        "Device_ID", "Hostname", "Platform_PID", "Base_PID", "Current_Version",
        "Target_Version", "Role", "Features_CSV", "Business_Criticality_1_5"
    ]
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"Inventory_Input is missing required columns: {', '.join(missing)}")

    devices: List[Device] = []
    for row in range(3, ws.max_row + 1):
        device_id = ws.cell(row, header["Device_ID"]).value
        if not device_id:
            continue
        platform_pid = str(ws.cell(row, header["Platform_PID"]).value or "").strip()
        base_pid = str(ws.cell(row, header["Base_PID"]).value or "").strip()
        family = optional_cell(ws, row, header, "Family") or derive_family(platform_pid, base_pid)
        site = optional_cell(ws, row, header, "Site")
        devices.append(
            Device(
                device_id=str(device_id).strip(),
                hostname=str(ws.cell(row, header["Hostname"]).value or "").strip(),
                platform_pid=platform_pid,
                base_pid=base_pid,
                current_version=str(ws.cell(row, header["Current_Version"]).value or "").strip(),
                target_version=str(ws.cell(row, header["Target_Version"]).value or "").strip(),
                role=str(ws.cell(row, header["Role"]).value or "").strip(),
                features=split_csvish(str(ws.cell(row, header["Features_CSV"]).value or "")),
                business_criticality=parse_int(str(ws.cell(row, header["Business_Criticality_1_5"]).value or "3")),
                site=site,
                family=norm(family),
            )
        )
    return devices


def load_bugs(ws) -> List[Bug]:
    header = get_header_map(ws)
    required = [
        "Bug_ID", "Headline", "Source_Type", "Product_Scope", "Affected_Releases",
        "Fixed_Releases", "Status", "Severity_Public", "Trigger_Features",
        "Trigger_Platforms", "Impact_Type", "Recommended_Action", "Source_URL"
    ]
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"Candidate_Bugs is missing required columns: {', '.join(missing)}")

    bugs: List[Bug] = []
    for row in range(3, ws.max_row + 1):
        bug_id = ws.cell(row, header["Bug_ID"]).value
        if not bug_id:
            continue
        bugs.append(
            Bug(
                bug_id=str(bug_id).strip(),
                headline=str(ws.cell(row, header["Headline"]).value or "").strip(),
                source_type=str(ws.cell(row, header["Source_Type"]).value or "").strip(),
                product_scope=str(ws.cell(row, header["Product_Scope"]).value or "").strip(),
                affected_releases=str(ws.cell(row, header["Affected_Releases"]).value or "").strip(),
                fixed_releases=str(ws.cell(row, header["Fixed_Releases"]).value or "").strip(),
                status=str(ws.cell(row, header["Status"]).value or "").strip(),
                severity_public=str(ws.cell(row, header["Severity_Public"]).value or "").strip(),
                trigger_features=split_csvish(str(ws.cell(row, header["Trigger_Features"]).value or "")),
                trigger_platforms=split_csvish(str(ws.cell(row, header["Trigger_Platforms"]).value or "")),
                impact_type=str(ws.cell(row, header["Impact_Type"]).value or "").strip(),
                recommended_action=str(ws.cell(row, header["Recommended_Action"]).value or "").strip(),
                source_url=str(ws.cell(row, header["Source_URL"]).value or "").strip(),
            )
        )
    return bugs


def score_bug(device: Device, bug: Bug) -> Optional[Dict[str, object]]:
    vmatch = version_in_text(device.current_version, bug.affected_releases)
    pmatch = platform_match(device.platform_pid, bug)
    fmatch = feature_match(device.features, bug)
    sev = severity_score(bug.severity_public)
    st = status_score(bug.status)

    score = 0
    if vmatch:
        score += 4
    elif extract_version_tokens(bug.affected_releases):
        score -= 1

    score += pmatch
    score += min(fmatch, 3)
    score += sev
    score += st
    score += max(0, device.business_criticality - 2)

    if not vmatch and score < 5:
        return None

    band = band_from_score(score)
    fixed_versions = ", ".join(extract_version_tokens(bug.fixed_releases))

    return {
        "Device_ID": device.device_id,
        "Hostname": device.hostname,
        "Site": device.site,
        "Family": device.family,
        "Role": device.role,
        "Platform_PID": device.platform_pid,
        "Current_Version": device.current_version,
        "Target_Version": device.target_version,
        "Bug_ID": bug.bug_id,
        "Headline": bug.headline,
        "Impact_Type": bug.impact_type,
        "Severity_Public": bug.severity_public,
        "Severity_Num": severity_numeric(bug.severity_public),
        "Status": bug.status,
        "Version_Match": "Yes" if vmatch else "No",
        "Platform_Match": pmatch,
        "Feature_Match_Count": fmatch,
        "Business_Criticality": device.business_criticality,
        "Score": score,
        "Risk_Band": band,
        "Fixed_Releases": bug.fixed_releases,
        "Fixed_Release_Tokens": fixed_versions,
        "Suggested_Action": action_hint(device, bug, score, vmatch),
        "Source_URL": bug.source_url,
    }


def clear_sheet(ws) -> None:
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws.sheet_view.showGridLines = False


def style_header_row(ws, row_num: int, headers: Sequence[str]) -> None:
    for idx, header in enumerate(headers, start=1):
        c = ws.cell(row_num, idx, header)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def apply_risk_fill(cell, band: str) -> None:
    fill = None
    if band == "High":
        fill = HIGH_FILL
    elif band == "Medium":
        fill = MED_FILL
    elif band == "Low":
        fill = LOW_FILL
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def write_table(ws, start_row: int, headers: Sequence[str], rows: List[Dict[str, object]], widths: Sequence[int], risk_band_header: Optional[str] = None) -> None:
    style_header_row(ws, start_row, headers)
    for r_idx, row in enumerate(rows, start=start_row + 1):
        for c_idx, header in enumerate(headers, start=1):
            cell = ws.cell(r_idx, c_idx, row.get(header))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if risk_band_header:
            risk_idx = headers.index(risk_band_header) + 1
            apply_risk_fill(ws.cell(r_idx, risk_idx), str(row.get(risk_band_header, "")))
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width


def write_auto_scored_sheet(wb, rows: List[Dict[str, object]]) -> None:
    ws = wb["Auto_Scored"] if "Auto_Scored" in wb.sheetnames else wb.create_sheet("Auto_Scored")
    clear_sheet(ws)
    ws.freeze_panes = "A3"
    ws["A1"] = "Auto scored output (generated by script v2)"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_FILL)

    headers = [
        "Device_ID", "Hostname", "Site", "Family", "Role", "Platform_PID", "Current_Version",
        "Target_Version", "Bug_ID", "Headline", "Impact_Type", "Severity_Public", "Status",
        "Version_Match", "Platform_Match", "Feature_Match_Count", "Business_Criticality",
        "Score", "Risk_Band", "Fixed_Releases", "Suggested_Action", "Source_URL"
    ]
    widths = [12, 18, 12, 12, 12, 18, 14, 14, 14, 40, 18, 12, 12, 12, 14, 18, 18, 10, 12, 22, 36, 50]
    write_table(ws, 2, headers, rows, widths, risk_band_header="Risk_Band")


def summarize_by_device(rows: List[Dict[str, object]], devices: List[Device]) -> List[DeviceSummary]:
    by_device: Dict[str, List[Dict[str, object]]] = {}
    for row in rows:
        by_device.setdefault(str(row["Device_ID"]), []).append(row)

    summaries: List[DeviceSummary] = []
    for device in devices:
        findings = by_device.get(device.device_id, [])
        findings_sorted = sorted(
            findings,
            key=lambda r: (-int(r["Score"]), int(r["Severity_Num"]), str(r["Bug_ID"]))
        )
        high_count = sum(1 for r in findings if r["Risk_Band"] == "High")
        medium_count = sum(1 for r in findings if r["Risk_Band"] == "Medium")
        low_count = sum(1 for r in findings if r["Risk_Band"] == "Low")
        max_score = max((int(r["Score"]) for r in findings), default=0)
        weighted = high_count * 10 + medium_count * 4 + low_count * 1 + max(0, device.business_criticality - 2)

        if findings_sorted:
            top = findings_sorted[0]
            top_bug_id = str(top["Bug_ID"])
            top_bug_headline = str(top["Headline"])
            top_action = str(top["Suggested_Action"])
        else:
            top_bug_id = ""
            top_bug_headline = ""
            top_action = "No matching bugs found in current candidate set."

        decision, reason = upgrade_decision_for_device(device, findings_sorted)

        summaries.append(
            DeviceSummary(
                device_id=device.device_id,
                hostname=device.hostname,
                site=device.site,
                family=device.family,
                role=device.role,
                platform_pid=device.platform_pid,
                current_version=device.current_version,
                target_version=device.target_version,
                business_criticality=device.business_criticality,
                findings_total=len(findings),
                high_count=high_count,
                medium_count=medium_count,
                low_count=low_count,
                max_score=max_score,
                weighted_risk_score=weighted,
                top_bug_id=top_bug_id,
                top_bug_headline=top_bug_headline,
                top_action=top_action,
                upgrade_decision=decision,
                decision_reason=reason,
            )
        )

    summaries.sort(
        key=lambda s: (-s.weighted_risk_score, -s.high_count, -s.max_score, s.hostname.lower())
    )
    return summaries


def upgrade_decision_for_device(device: Device, findings_sorted: List[Dict[str, object]]) -> Tuple[str, str]:
    if not findings_sorted:
        return ("GO", "No candidate bugs matched this node in the current scrub dataset.")

    current_target_same = norm(device.current_version) == norm(device.target_version)
    high_open_version_matches = [
        r for r in findings_sorted
        if r["Risk_Band"] == "High"
        and r["Version_Match"] == "Yes"
        and "open" in norm(str(r["Status"]))
    ]
    severe_hits = [
        r for r in findings_sorted
        if int(r["Severity_Num"]) <= 2 and r["Version_Match"] == "Yes"
    ]
    medium_hits = [r for r in findings_sorted if r["Risk_Band"] == "Medium"]

    if high_open_version_matches:
        top = high_open_version_matches[0]
        return (
            "NO-GO",
            f"High-risk open bug {top['Bug_ID']} matches the current release; validate workaround or move target version before change.",
        )

    if severe_hits and current_target_same:
        top = severe_hits[0]
        return (
            "NO-GO",
            f"Current and target releases are the same while severe bug {top['Bug_ID']} still matches this node.",
        )

    if severe_hits or len(medium_hits) >= 3:
        top = severe_hits[0] if severe_hits else medium_hits[0]
        return (
            "CONDITIONAL",
            f"Proceed only with mitigation and validation; bug {top['Bug_ID']} is still relevant for this node.",
        )

    return (
        "GO",
        "No blocking high-risk conditions identified in the current candidate set; continue with standard validation.",
    )


def update_summary_sheet(wb, scored_rows: List[Dict[str, object]], device_summaries: List[DeviceSummary], filters_text: str) -> None:
    ws = wb["Exec_Summary_Auto"] if "Exec_Summary_Auto" in wb.sheetnames else wb.create_sheet("Exec_Summary_Auto")
    clear_sheet(ws)
    ws["A1"] = "Auto summary"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_FILL)

    high = sum(1 for r in scored_rows if r["Risk_Band"] == "High")
    med = sum(1 for r in scored_rows if r["Risk_Band"] == "Medium")
    low = sum(1 for r in scored_rows if r["Risk_Band"] == "Low")
    go_count = sum(1 for d in device_summaries if d.upgrade_decision == "GO")
    cond_count = sum(1 for d in device_summaries if d.upgrade_decision == "CONDITIONAL")
    nogo_count = sum(1 for d in device_summaries if d.upgrade_decision == "NO-GO")

    metrics = [
        ("Applied filters", filters_text or "None"),
        ("Devices in scope", len(device_summaries)),
        ("Auto-scored findings", len(scored_rows)),
        ("High", high),
        ("Medium", med),
        ("Low", low),
        ("GO", go_count),
        ("CONDITIONAL", cond_count),
        ("NO-GO", nogo_count),
    ]

    for idx, (label, value) in enumerate(metrics, start=3):
        ws[f"A{idx}"] = label
        ws[f"B{idx}"] = value
        ws[f"A{idx}"].font = Font(bold=True)
        ws[f"B{idx}"].alignment = Alignment(wrap_text=True)

    ws["D3"] = "Top 5 nodes by weighted risk"
    ws["D3"].font = Font(bold=True, color="FFFFFF")
    ws["D3"].fill = PatternFill("solid", fgColor=HEADER_FILL)
    top_headers = ["Hostname", "Site", "Role", "Weighted_Risk_Score", "Decision"]
    style_header_row(ws, 4, top_headers)
    for i, summary in enumerate(device_summaries[:5], start=5):
        values = [summary.hostname, summary.site, summary.role, summary.weighted_risk_score, summary.upgrade_decision]
        for j, value in enumerate(values, start=4):
            ws.cell(i, j, value).alignment = Alignment(wrap_text=True)
        apply_risk_fill(ws.cell(i, 8), "High" if summary.upgrade_decision == "NO-GO" else "Medium" if summary.upgrade_decision == "CONDITIONAL" else "Low")

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 32
    for col in ["D", "E", "F", "G", "H"]:
        ws.column_dimensions[col].width = 18


def write_device_ranking_sheet(wb, device_summaries: List[DeviceSummary]) -> None:
    ws = wb["Device_Ranking"] if "Device_Ranking" in wb.sheetnames else wb.create_sheet("Device_Ranking")
    clear_sheet(ws)
    ws.freeze_panes = "A3"
    ws["A1"] = "Per-device ranking"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_FILL)

    headers = [
        "Rank", "Device_ID", "Hostname", "Site", "Family", "Role", "Platform_PID",
        "Current_Version", "Target_Version", "Business_Criticality", "Findings_Total",
        "High_Count", "Medium_Count", "Low_Count", "Max_Score", "Weighted_Risk_Score",
        "Top_Bug_ID", "Top_Bug_Headline", "Top_Action", "Upgrade_Decision", "Decision_Reason"
    ]
    rows = []
    for rank, s in enumerate(device_summaries, start=1):
        rows.append({
            "Rank": rank,
            "Device_ID": s.device_id,
            "Hostname": s.hostname,
            "Site": s.site,
            "Family": s.family,
            "Role": s.role,
            "Platform_PID": s.platform_pid,
            "Current_Version": s.current_version,
            "Target_Version": s.target_version,
            "Business_Criticality": s.business_criticality,
            "Findings_Total": s.findings_total,
            "High_Count": s.high_count,
            "Medium_Count": s.medium_count,
            "Low_Count": s.low_count,
            "Max_Score": s.max_score,
            "Weighted_Risk_Score": s.weighted_risk_score,
            "Top_Bug_ID": s.top_bug_id,
            "Top_Bug_Headline": s.top_bug_headline,
            "Top_Action": s.top_action,
            "Upgrade_Decision": s.upgrade_decision,
            "Decision_Reason": s.decision_reason,
        })
    widths = [8, 12, 18, 12, 12, 12, 18, 14, 14, 18, 14, 10, 12, 10, 10, 18, 14, 36, 36, 16, 44]
    write_table(ws, 2, headers, rows, widths)
    decision_idx = headers.index("Upgrade_Decision") + 1
    for r in range(3, ws.max_row + 1):
        decision = str(ws.cell(r, decision_idx).value or "")
        band = "High" if decision == "NO-GO" else "Medium" if decision == "CONDITIONAL" else "Low"
        apply_risk_fill(ws.cell(r, decision_idx), band)


def write_upgrade_decision_sheet(wb, device_summaries: List[DeviceSummary]) -> None:
    ws = wb["Upgrade_Decision"] if "Upgrade_Decision" in wb.sheetnames else wb.create_sheet("Upgrade_Decision")
    clear_sheet(ws)
    ws.freeze_panes = "A3"
    ws["A1"] = "Go / No-Go recommendation by node"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=HEADER_FILL)

    headers = [
        "Hostname", "Site", "Role", "Family", "Current_Version", "Target_Version",
        "Upgrade_Decision", "Decision_Reason", "Top_Bug_ID", "Top_Bug_Headline", "Top_Action"
    ]
    rows = []
    for s in device_summaries:
        rows.append({
            "Hostname": s.hostname,
            "Site": s.site,
            "Role": s.role,
            "Family": s.family,
            "Current_Version": s.current_version,
            "Target_Version": s.target_version,
            "Upgrade_Decision": s.upgrade_decision,
            "Decision_Reason": s.decision_reason,
            "Top_Bug_ID": s.top_bug_id,
            "Top_Bug_Headline": s.top_bug_headline,
            "Top_Action": s.top_action,
        })
    widths = [18, 12, 12, 12, 14, 14, 16, 44, 14, 36, 36]
    write_table(ws, 2, headers, rows, widths)
    decision_idx = headers.index("Upgrade_Decision") + 1
    for r in range(3, ws.max_row + 1):
        decision = str(ws.cell(r, decision_idx).value or "")
        band = "High" if decision == "NO-GO" else "Medium" if decision == "CONDITIONAL" else "Low"
        apply_risk_fill(ws.cell(r, decision_idx), band)


def parse_filter_values(raw: Optional[str]) -> Optional[set[str]]:
    if not raw:
        return None
    values = {norm(part) for part in re.split(r"[;,]", raw) if norm(part)}
    return values or None


def filter_devices(devices: List[Device], site: Optional[str], role: Optional[str], family: Optional[str]) -> List[Device]:
    site_values = parse_filter_values(site)
    role_values = parse_filter_values(role)
    family_values = parse_filter_values(family)

    def keep(device: Device) -> bool:
        if site_values and norm(device.site) not in site_values:
            return False
        if role_values and norm(device.role) not in role_values:
            return False
        if family_values and norm(device.family) not in family_values:
            return False
        return True

    return [d for d in devices if keep(d)]


def filter_description(site: Optional[str], role: Optional[str], family: Optional[str]) -> str:
    parts = []
    if site:
        parts.append(f"site={site}")
    if role:
        parts.append(f"role={role}")
    if family:
        parts.append(f"family={family}")
    return ", ".join(parts)


def export_csv(path: Path, rows: List[Dict[str, object]], headers: Sequence[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(headers))
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in headers})


def workbook_mode(
    workbook_path: Path,
    output_path: Path,
    site: Optional[str] = None,
    role: Optional[str] = None,
    family: Optional[str] = None,
    export_scored_csv: Optional[Path] = None,
    export_device_csv: Optional[Path] = None,
) -> int:
    wb = load_workbook(workbook_path)
    if "Inventory_Input" not in wb.sheetnames or "Candidate_Bugs" not in wb.sheetnames:
        print("Workbook must contain sheets Inventory_Input and Candidate_Bugs.", file=sys.stderr)
        return 2

    try:
        devices = load_devices(wb["Inventory_Input"])
        bugs = load_bugs(wb["Candidate_Bugs"])
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        return 2

    devices = filter_devices(devices, site=site, role=role, family=family)
    if not devices:
        print("No devices remain after filters. Check --site, --role, and --family values.", file=sys.stderr)
        return 2

    scored_rows: List[Dict[str, object]] = []
    for device in devices:
        for bug in bugs:
            row = score_bug(device, bug)
            if row:
                scored_rows.append(row)

    scored_rows.sort(key=lambda r: (-int(r["Score"]), str(r["Device_ID"]), str(r["Bug_ID"])))
    device_summaries = summarize_by_device(scored_rows, devices)
    filters_text = filter_description(site, role, family)

    write_auto_scored_sheet(wb, scored_rows)
    write_device_ranking_sheet(wb, device_summaries)
    write_upgrade_decision_sheet(wb, device_summaries)
    update_summary_sheet(wb, scored_rows, device_summaries, filters_text)
    wb.save(output_path)

    if export_scored_csv:
        headers = [
            "Device_ID", "Hostname", "Site", "Family", "Role", "Platform_PID", "Current_Version",
            "Target_Version", "Bug_ID", "Headline", "Impact_Type", "Severity_Public", "Status",
            "Version_Match", "Platform_Match", "Feature_Match_Count", "Business_Criticality",
            "Score", "Risk_Band", "Fixed_Releases", "Suggested_Action", "Source_URL"
        ]
        export_csv(export_scored_csv, scored_rows, headers)

    if export_device_csv:
        device_rows = [{
            "Device_ID": s.device_id,
            "Hostname": s.hostname,
            "Site": s.site,
            "Family": s.family,
            "Role": s.role,
            "Platform_PID": s.platform_pid,
            "Current_Version": s.current_version,
            "Target_Version": s.target_version,
            "Business_Criticality": s.business_criticality,
            "Findings_Total": s.findings_total,
            "High_Count": s.high_count,
            "Medium_Count": s.medium_count,
            "Low_Count": s.low_count,
            "Max_Score": s.max_score,
            "Weighted_Risk_Score": s.weighted_risk_score,
            "Top_Bug_ID": s.top_bug_id,
            "Top_Bug_Headline": s.top_bug_headline,
            "Top_Action": s.top_action,
            "Upgrade_Decision": s.upgrade_decision,
            "Decision_Reason": s.decision_reason,
        } for s in device_summaries]
        headers = list(device_rows[0].keys()) if device_rows else []
        export_csv(export_device_csv, device_rows, headers)

    print(f"Devices in scope: {len(devices)}")
    print(f"Wrote {len(scored_rows)} findings to {output_path}")
    print(f"GO / CONDITIONAL / NO-GO: {sum(1 for s in device_summaries if s.upgrade_decision == 'GO')} / {sum(1 for s in device_summaries if s.upgrade_decision == 'CONDITIONAL')} / {sum(1 for s in device_summaries if s.upgrade_decision == 'NO-GO')}")
    return 0


def get_token(client_id: str, client_secret: str) -> str:
    response = requests.post(
        TOKEN_URL,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        data={
            "grant_type": "client_credentials",
            "client_id": client_id,
            "client_secret": client_secret,
        },
        timeout=30,
    )
    response.raise_for_status()
    return response.json()["access_token"]


def pull_bugs(product_series: str, affected_release: str, status: Optional[str], modified_date: str = "5", page_index: int = 1) -> dict:
    client_id = os.getenv("CISCO_CLIENT_ID")
    client_secret = os.getenv("CISCO_CLIENT_SECRET")
    if not client_id or not client_secret:
        raise RuntimeError("Set CISCO_CLIENT_ID and CISCO_CLIENT_SECRET in the environment.")

    token = get_token(client_id, client_secret)
    url = (
        f"{BUG_API_BASE}/product_series/{quote(product_series, safe='')}"
        f"/affected_releases/{quote(affected_release, safe='()')}"
    )
    params = {"modified_date": modified_date, "page_index": page_index, "sort_by": "modified_date"}
    if status:
        params["status"] = status
    response = requests.get(
        url,
        headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
        params=params,
        timeout=60,
    )
    response.raise_for_status()
    return response.json()


def api_mode(product_series: str, affected_release: str, output_csv: Path, status: Optional[str]) -> int:
    payload = pull_bugs(product_series=product_series, affected_release=affected_release, status=status)
    bugs = payload.get("bugs", [])
    headers = [
        "bug_id", "headline", "severity", "status", "last_modified_date",
        "known_affected_releases", "known_fixed_releases", "support_case_count",
        "product", "source_url"
    ]
    with output_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for bug in bugs:
            bug_id = bug.get("bug_id", "")
            writer.writerow(
                {
                    "bug_id": bug_id,
                    "headline": bug.get("headline", ""),
                    "severity": bug.get("severity", ""),
                    "status": bug.get("status", ""),
                    "last_modified_date": bug.get("last_modified_date", ""),
                    "known_affected_releases": bug.get("known_affected_releases", ""),
                    "known_fixed_releases": bug.get("known_fixed_releases", ""),
                    "support_case_count": bug.get("support_case_count", ""),
                    "product": bug.get("product", "") or bug.get("product_series", ""),
                    "source_url": f"https://bst.cisco.com/quickview/bug/{bug_id}" if bug_id else "",
                }
            )
    print(f"Wrote {len(bugs)} API rows to {output_csv}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Cisco bug scrub helper v2")
    parser.add_argument("--workbook", type=Path, help="Input workbook containing Inventory_Input and Candidate_Bugs sheets")
    parser.add_argument("--output", type=Path, help="Output workbook path")
    parser.add_argument("--site", help="Filter devices by Site column; accepts comma-separated values")
    parser.add_argument("--role", help="Filter devices by Role column; accepts comma-separated values")
    parser.add_argument("--family", help="Filter devices by Family column or derived family; accepts comma-separated values")
    parser.add_argument("--export-scored-csv", type=Path, help="Optional CSV export of Auto_Scored")
    parser.add_argument("--export-device-csv", type=Path, help="Optional CSV export of per-device summary")
    parser.add_argument("--api-product-series", help='Cisco product series, e.g. "Cisco Nexus 9000 Series Switches"')
    parser.add_argument("--api-affected-release", help='Affected release, e.g. "10.4(4)M"')
    parser.add_argument("--api-output", type=Path, help="CSV path for API pull output")
    parser.add_argument("--api-status", choices=["O", "F", "T"], help="Optional API bug status filter")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    if args.api_product_series and args.api_affected_release and args.api_output:
        return api_mode(
            product_series=args.api_product_series,
            affected_release=args.api_affected_release,
            output_csv=args.api_output,
            status=args.api_status,
        )

    if args.workbook and args.output:
        return workbook_mode(
            workbook_path=args.workbook,
            output_path=args.output,
            site=args.site,
            role=args.role,
            family=args.family,
            export_scored_csv=args.export_scored_csv,
            export_device_csv=args.export_device_csv,
        )

    parser.print_help()
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
