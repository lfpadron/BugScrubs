from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook

from bugscrub.normalization.models import DeviceRecord, InventoryBatch, InventoryRowRecord, InventoryWorkbook


HEADER_ALIASES = {
    "hostname": {"hostname", "host name", "device name", "device_name", "switch name", "switch", "node name"},
    "model": {"model", "model number", "platform", "platform model", "chassis", "product model"},
    "pid": {"pid", "platform pid", "product id", "product_id", "platform_id", "base pid", "sku"},
    "serial": {"serial", "serial number", "sn", "chassis serial", "system serial number"},
    "os_version": {
        "os version",
        "os_version",
        "version",
        "current version",
        "current_version",
        "software version",
        "nx-os version",
        "ios xe version",
    },
    "target_version": {"target version", "target_version", "planned version", "recommended version"},
    "site": {"site", "location", "campus", "datacenter"},
    "role": {"role", "device role", "node role"},
    "family": {"family", "platform family", "product family"},
    "features": {"features", "enabled features", "feature set", "trigger features"},
    "business_criticality": {"business criticality", "criticality", "business_criticality", "priority"},
}

GENERATED_INVENTORY_HEADERS = [
    "Device Name",
    "Model Number",
    "Platform PID",
    "Serial Number",
    "Current Version",
    "Target Version",
    "Site",
    "Role",
    "Product Family",
    "Enabled Features",
    "Criticality",
]


class Normalizer:
    """Converts parser output into canonical inventory records."""

    def normalize_parsed_device(
        self,
        *,
        parsed_record: dict[str, Any],
        platform_family: str,
        support_level: str,
        session_dir: Path,
        inventory_path: Path | None,
        source_count: int,
        warnings: list[str] | None = None,
    ) -> InventoryBatch:
        return self.normalize_parsed_devices(
            parsed_records=[parsed_record],
            platform_family=platform_family,
            support_level=support_level,
            session_dir=session_dir,
            inventory_path=inventory_path,
            source_count=source_count,
            warnings=warnings,
        )

    def normalize_parsed_devices(
        self,
        *,
        parsed_records: list[dict[str, Any]],
        platform_family: str,
        support_level: str,
        session_dir: Path,
        inventory_path: Path | None,
        source_count: int,
        warnings: list[str] | None = None,
    ) -> InventoryBatch:
        workbook, inventory_rows, workbook_warnings = parse_inventory_workbook(
            inventory_path=inventory_path,
            session_id=session_dir.name,
        )
        warning_list = list(warnings or [])
        if inventory_path is not None and not workbook.file_name:
            warning_list.append("Inventory workbook metadata could not be loaded.")
        warning_list.extend(workbook_warnings)

        platform_key = platform_family.lower()
        devices = [
            device
            for parsed_record in parsed_records
            if (device := self._build_device_record(
                parsed_record=parsed_record,
                platform_family=platform_key,
                support_level=support_level,
                session_id=session_dir.name,
                workbook=workbook,
            )) is not None
        ]
        if parsed_records and not devices:
            warning_list.append("No parsed devices produced usable normalized records.")
        return InventoryBatch(
            session_id=session_dir.name,
            platform_family=platform_key,
            support_level=support_level,
            session_dir=str(session_dir),
            inventory_workbook=workbook,
            inventory_rows=inventory_rows,
            devices=devices,
            source_count=source_count,
            warnings=warning_list,
        )

    def _build_device_record(
        self,
        *,
        parsed_record: dict[str, Any],
        platform_family: str,
        support_level: str,
        session_id: str,
        workbook: InventoryWorkbook,
    ) -> DeviceRecord | None:
        version = str(parsed_record.get("nxos_version") or parsed_record.get("os_version") or "")
        hostname = str(parsed_record.get("device_name") or parsed_record.get("bundle_name") or "")
        model = str(parsed_record.get("model") or "")
        pid = str(parsed_record.get("pid") or "")
        serial = str(parsed_record.get("serial") or "")
        features = [str(feature) for feature in parsed_record.get("features", [])]

        if not any([hostname, model, pid, serial, version, features]):
            return None

        return DeviceRecord(
            session_id=session_id,
            hostname=hostname,
            vendor="Cisco",
            platform_family=platform_family,
            support_level=support_level,
            model=model,
            pid=pid,
            serial=serial,
            os_name="NX-OS" if platform_family == "nexus" else "IOS XE",
            os_version=version,
            inventory_file=workbook.file_name,
            inventory_sheet_names=list(workbook.sheet_names),
            features=features,
        )


def generate_inventory_workbook_from_parsed_records(
    *,
    parsed_records: list[dict[str, Any]],
    destination: Path,
    platform_family: str,
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Inventory_Input"
    worksheet.append(GENERATED_INVENTORY_HEADERS)

    generated_row_count = 0
    for parsed_record in parsed_records:
        row = build_generated_inventory_row(
            parsed_record=parsed_record,
            platform_family=platform_family,
        )
        if row is None:
            continue
        worksheet.append(row)
        generated_row_count += 1

    workbook.create_sheet("Candidate_Bugs")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()

    if generated_row_count == 0:
        raise ValueError("No parsed devices were available to generate an inventory workbook.")

    return destination


def parse_inventory_workbook(
    inventory_path: Path | None,
    *,
    session_id: str,
) -> tuple[InventoryWorkbook, list[InventoryRowRecord], list[str]]:
    if inventory_path is None or not inventory_path.exists():
        return InventoryWorkbook(), [], []

    try:
        workbook = load_workbook(inventory_path, read_only=True, data_only=True)
    except Exception:
        return InventoryWorkbook(), [], [f"Inventory workbook `{inventory_path.name}` could not be opened or is corrupt."]

    try:
        sheet_names = list(workbook.sheetnames)
        selected_sheet, normalized_columns, header_row_index = select_inventory_sheet(workbook)
        if selected_sheet is None:
            return (
                InventoryWorkbook(
                    file_name=inventory_path.name,
                    sheet_names=sheet_names,
                    sheet_count=len(sheet_names),
                ),
                [],
                ["Inventory workbook did not contain a recognizable header row."],
            )

        rows = parse_inventory_rows(
            worksheet=workbook[selected_sheet],
            session_id=session_id,
            normalized_columns=normalized_columns,
            header_row_index=header_row_index,
        )
        return (
            InventoryWorkbook(
                file_name=inventory_path.name,
                sheet_names=sheet_names,
                sheet_count=len(sheet_names),
                selected_sheet=selected_sheet,
                normalized_columns=[column for column in normalized_columns if column],
                parsed_row_count=len(rows),
            ),
            rows,
            [],
        )
    finally:
        workbook.close()


def inspect_inventory_workbook(inventory_path: Path | None) -> InventoryWorkbook:
    workbook, _, _ = parse_inventory_workbook(inventory_path=inventory_path, session_id="")
    return workbook


def select_inventory_sheet(workbook: Any) -> tuple[str | None, list[str], int]:
    preferred_names = ("inventory_input", "inventory", "devices")

    for sheet_name in workbook.sheetnames:
        if normalize_header_name(sheet_name) in preferred_names:
            normalized_columns, header_row_index = detect_header_row(workbook[sheet_name])
            if normalized_columns:
                return sheet_name, normalized_columns, header_row_index

    for sheet_name in workbook.sheetnames:
        normalized_columns, header_row_index = detect_header_row(workbook[sheet_name])
        if normalized_columns:
            return sheet_name, normalized_columns, header_row_index

    return None, [], 0


def detect_header_row(worksheet: Any) -> tuple[list[str], int]:
    for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = [stringify_cell(cell) for cell in row]
        normalized_columns = [map_header(value) if value else "" for value in values]
        recognized_columns = [column for column in normalized_columns if column]
        if len(recognized_columns) >= 2 and "hostname" in recognized_columns:
            return normalized_columns, index
    return [], 0


def parse_inventory_rows(
    *,
    worksheet: Any,
    session_id: str,
    normalized_columns: list[str],
    header_row_index: int,
) -> list[InventoryRowRecord]:
    rows: list[InventoryRowRecord] = []
    data_rows = worksheet.iter_rows(min_row=header_row_index + 1, values_only=True)
    for offset, row in enumerate(data_rows, start=header_row_index + 1):
        values = [stringify_cell(cell) for cell in row]
        if not any(values):
            continue

        mapped = {
            column: values[index] if index < len(values) else ""
            for index, column in enumerate(normalized_columns)
            if column
        }
        row_record = InventoryRowRecord(
            session_id=session_id,
            sheet_name=worksheet.title,
            row_number=offset,
            hostname=mapped.get("hostname", ""),
            model=mapped.get("model", ""),
            pid=mapped.get("pid", ""),
            serial=mapped.get("serial", ""),
            os_version=mapped.get("os_version", ""),
            target_version=mapped.get("target_version", ""),
            site=mapped.get("site", ""),
            role=mapped.get("role", ""),
            family=mapped.get("family", ""),
            features=parse_feature_list(mapped.get("features", "")),
            business_criticality=parse_optional_int(mapped.get("business_criticality", "")),
        )
        if not row_record.hostname and not any(
            [row_record.model, row_record.pid, row_record.serial, row_record.os_version]
        ):
            continue
        rows.append(row_record)
    return rows


def map_header(value: str) -> str:
    normalized = normalize_header_name(value)
    for canonical_name, aliases in HEADER_ALIASES.items():
        if normalized in {normalize_header_name(alias) for alias in aliases}:
            return canonical_name
    return ""


def normalize_header_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def stringify_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_feature_list(value: str) -> list[str]:
    if not value:
        return []
    parts = re.split(r"[;,/]|(?:\s+\+\s+)", value)
    return [part.strip() for part in parts if part and part.strip()]


def parse_optional_int(value: str) -> int | None:
    try:
        return int(str(value).strip())
    except Exception:
        return None


def build_generated_inventory_row(
    *,
    parsed_record: dict[str, Any],
    platform_family: str,
) -> list[object] | None:
    hostname = str(parsed_record.get("device_name") or parsed_record.get("bundle_name") or "").strip()
    model = str(parsed_record.get("model") or "").strip()
    pid = str(parsed_record.get("pid") or "").strip()
    serial = str(parsed_record.get("serial") or "").strip()
    os_version = str(parsed_record.get("nxos_version") or parsed_record.get("os_version") or "").strip()
    features = normalize_feature_values(parsed_record.get("features", []))
    family = infer_generated_family(platform_family=platform_family, pid=pid, model=model)

    if not any([hostname, model, pid, serial, os_version, features]):
        return None

    return [
        hostname,
        model,
        pid,
        serial,
        os_version,
        "",
        "",
        "",
        family,
        "; ".join(features),
        "",
    ]


def normalize_feature_values(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str):
        return parse_feature_list(value)
    return []


def infer_generated_family(*, platform_family: str, pid: str, model: str) -> str:
    combined = f"{pid} {model}".lower()
    if platform_family == "nexus":
        if "n9k" in combined or "c93" in combined or "c95" in combined:
            return "nexus9000"
        if "n7k" in combined:
            return "nexus7000"
        return "nexus"

    if platform_family == "catalyst":
        for family in ("c9300", "c9500", "c9200", "c9400", "c9600"):
            if family in combined:
                return f"catalyst{family[1:]}"
        return "catalyst"

    return platform_family
