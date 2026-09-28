from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet


DISCREPANCY_COLUMNS = [
    "pair_id",
    "row_source",
    "discrepancy_status",
    "discrepancy_summary",
    "hostname",
    "model",
    "pid",
    "serial",
    "os_version",
    "target_version",
    "site",
    "role",
    "family",
    "features",
    "discrepancy_fields",
]

DISCREPANCY_HEADERS = {
    "pair_id": "Pair ID",
    "row_source": "Row Source",
    "discrepancy_status": "Status",
    "discrepancy_summary": "Summary",
    "hostname": "Hostname",
    "model": "Model",
    "pid": "PID",
    "serial": "Serial",
    "os_version": "Version",
    "target_version": "Target Version",
    "site": "Site",
    "role": "Role",
    "family": "Family",
    "features": "Features",
    "discrepancy_fields": "Differing Fields",
}

BUG_FINDING_COLUMNS = [
    "hostname",
    "platform_family",
    "bug_id",
    "headline",
    "severity",
    "score",
    "current_version",
    "target_version",
    "remediation_status",
    "version_match_type",
    "matched_release",
    "matched_on",
    "fixed_releases",
    "recommended_action",
    "rationale",
]

BUG_FINDING_HEADERS = {
    "hostname": "Hostname",
    "platform_family": "Platform Family",
    "bug_id": "Bug ID",
    "headline": "Headline",
    "severity": "Severity",
    "score": "Score",
    "current_version": "Current Version",
    "target_version": "Target Version",
    "remediation_status": "Remediation Status",
    "version_match_type": "Version Match Type",
    "matched_release": "Matched Release",
    "matched_on": "Matched On",
    "fixed_releases": "Fixed Releases",
    "recommended_action": "Recommended Action",
    "rationale": "Rationale",
}


class ExcelExporter:
    """Operational Excel exporter for discrepancy outputs."""

    def export_discrepancies(
        self,
        discrepancy_rows: list[dict[str, object]],
        destination: Path,
        bug_findings: list[dict[str, object]] | None = None,
    ) -> Path:
        workbook = self.build_discrepancy_workbook(discrepancy_rows, bug_findings=bug_findings or [])
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        workbook.close()
        return destination

    def export_discrepancies_to_bytes(
        self,
        discrepancy_rows: list[dict[str, object]],
        bug_findings: list[dict[str, object]] | None = None,
    ) -> bytes:
        workbook = self.build_discrepancy_workbook(discrepancy_rows, bug_findings=bug_findings or [])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()

    def build_discrepancy_workbook(
        self,
        discrepancy_rows: list[dict[str, object]],
        *,
        bug_findings: list[dict[str, object]],
    ) -> Workbook:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Discrepancias"

        self._write_header(worksheet, DISCREPANCY_COLUMNS, DISCREPANCY_HEADERS)
        for row_index, row in enumerate(self._order_rows(discrepancy_rows), start=2):
            for column_index, column_name in enumerate(DISCREPANCY_COLUMNS, start=1):
                worksheet.cell(row=row_index, column=column_index, value=self._format_cell(row.get(column_name)))

        self._autosize_columns(worksheet)
        worksheet.freeze_panes = "A2"

        findings_sheet = workbook.create_sheet("Bug Findings")
        self._write_header(findings_sheet, BUG_FINDING_COLUMNS, BUG_FINDING_HEADERS)
        for row_index, row in enumerate(self._order_bug_findings(bug_findings), start=2):
            for column_index, column_name in enumerate(BUG_FINDING_COLUMNS, start=1):
                findings_sheet.cell(row=row_index, column=column_index, value=self._format_cell(row.get(column_name)))

        self._autosize_columns(findings_sheet)
        findings_sheet.freeze_panes = "A2"
        return workbook

    def read_export_rows(self, export_path: Path, sheet_name: str = "Discrepancias") -> list[tuple[object, ...]]:
        workbook = load_workbook(export_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet_name]
            return list(worksheet.iter_rows(min_row=2, values_only=True))
        finally:
            workbook.close()

    def _write_header(self, worksheet: Worksheet, columns: list[str], headers: dict[str, str]) -> None:
        for column_index, column_name in enumerate(columns, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=headers[column_name])
            cell.font = Font(bold=True)

    def _order_rows(self, discrepancy_rows: list[dict[str, object]]) -> list[dict[str, object]]:
        source_order = {"cliente": 0, "descubierto": 1}
        return sorted(
            discrepancy_rows,
            key=lambda row: (
                str(row.get("pair_id", "")),
                source_order.get(str(row.get("row_source", "")), 99),
            ),
        )

    def _format_cell(self, value: object) -> object:
        if isinstance(value, list):
            return ", ".join(str(item) for item in value)
        return value

    def _order_bug_findings(self, bug_findings: list[dict[str, object]]) -> list[dict[str, object]]:
        return sorted(
            bug_findings,
            key=lambda row: (
                str(row.get("hostname", "")),
                -int(row.get("score", 0) or 0),
                str(row.get("bug_id", "")),
            ),
        )

    def _autosize_columns(self, worksheet: Worksheet) -> None:
        for column_cells in worksheet.columns:
            values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
            max_length = max((len(value) for value in values), default=0)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 48)
