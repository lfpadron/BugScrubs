from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

from bugscrub.exporters.pareto import ParetoChartAsset, build_pareto_chart_asset, format_filter_lines
from bugscrub.exporters.risk_charts import (
    DashboardChartAsset,
    build_bug_breakdown_chart_asset,
    build_platform_breakdown_chart_asset,
    build_treemap_chart_asset,
)


SLIDE_WIDTH = Emu(9_144_000)
SLIDE_HEIGHT = Emu(5_143_500)

TITLE_X = Emu(457_200)
TITLE_Y = Emu(228_600)
TITLE_CX = Emu(8_230_000)
TITLE_CY = Emu(548_640)

BODY_X = Emu(457_200)
BODY_Y = Emu(914_400)
BODY_CX = Emu(2_743_200)
BODY_CY = Emu(3_657_600)

IMAGE_X = Emu(3_474_720)
IMAGE_Y = Emu(914_400)
IMAGE_CX = Emu(5_212_080)
IMAGE_CY = Emu(3_657_600)

TITLE_FONT_SIZE = Pt(24)
BODY_FONT_SIZE = Pt(13)

ChartAsset = DashboardChartAsset | ParetoChartAsset


class PowerPointExporter:
    """Executive PowerPoint exporter backed by python-pptx for desktop compatibility."""

    def export_pptx(
        self,
        *,
        session_id: str,
        platform_family: str,
        dashboard_data: dict[str, object],
        destination: Path,
        filter_summary: dict[str, object] | None = None,
    ) -> Path:
        pptx_bytes = self.export_pptx_to_bytes(
            session_id=session_id,
            platform_family=platform_family,
            dashboard_data=dashboard_data,
            filter_summary=filter_summary or {},
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(pptx_bytes)
        return destination

    def export_pptx_to_bytes(
        self,
        *,
        session_id: str,
        platform_family: str,
        dashboard_data: dict[str, object],
        filter_summary: dict[str, object] | None = None,
    ) -> bytes:
        filter_summary = filter_summary or {}
        metrics = dict(dashboard_data.get("metrics", {}))
        narrative_lines = list(dashboard_data.get("narrative_lines", []))
        pareto_quick_analysis_rows = list(dashboard_data.get("pareto_quick_analysis_rows", []))
        chart_assets = self._build_chart_assets(dashboard_data=dashboard_data, filter_summary=filter_summary)

        slides: list[dict[str, object]] = [
            {
                "title": "BugScrub Executive Report",
                "summary_lines": self._summary_lines(
                    session_id=session_id,
                    platform_family=platform_family,
                    metrics=metrics,
                    filter_summary=filter_summary,
                    narrative_lines=narrative_lines,
                    pareto_quick_analysis_rows=pareto_quick_analysis_rows,
                ),
                "asset": None,
            }
        ]
        for asset in chart_assets:
            slides.append(
                {
                    "title": asset.title,
                    "summary_lines": self._chart_summary_lines(filter_summary),
                    "asset": asset,
                }
            )

        presentation = Presentation()
        presentation.slide_width = SLIDE_WIDTH
        presentation.slide_height = SLIDE_HEIGHT

        for slide_definition in slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            self._add_textbox(
                slide=slide,
                x=TITLE_X,
                y=TITLE_Y,
                cx=TITLE_CX,
                cy=TITLE_CY,
                lines=[str(slide_definition["title"])],
                font_size=TITLE_FONT_SIZE,
                bold=True,
            )
            self._add_textbox(
                slide=slide,
                x=BODY_X,
                y=BODY_Y,
                cx=BODY_CX,
                cy=BODY_CY,
                lines=[str(line) for line in list(slide_definition["summary_lines"])],
                font_size=BODY_FONT_SIZE,
                bold=False,
            )
            asset = slide_definition["asset"]
            if asset is not None:
                slide.shapes.add_picture(
                    BytesIO(asset.jpeg_bytes),
                    IMAGE_X,
                    IMAGE_Y,
                    width=IMAGE_CX,
                    height=IMAGE_CY,
                )

        buffer = BytesIO()
        presentation.save(buffer)
        return buffer.getvalue()

    def _build_chart_assets(
        self,
        *,
        dashboard_data: dict[str, object],
        filter_summary: dict[str, object],
    ) -> list[ChartAsset]:
        return [
            build_bug_breakdown_chart_asset(
                stack_rows=list(dashboard_data.get("bug_severity_stack_rows", [])),
                filter_summary=filter_summary,
                title="Stacked Breakdown by Bug",
            ),
            build_platform_breakdown_chart_asset(
                stack_rows=list(dashboard_data.get("platform_stack_rows", [])),
                filter_summary=filter_summary,
                title="Stacked Breakdown by Platform",
            ),
            build_treemap_chart_asset(
                treemap_rows=list(dashboard_data.get("treemap_rows", [])),
                filter_summary=filter_summary,
                title="Treemap of Prioritized Findings",
            ),
            build_pareto_chart_asset(
                pareto_rows=list(dashboard_data.get("pareto_rows", [])),
                filter_summary=filter_summary,
                title="Pareto of Bug Remediation Impact",
            ),
        ]

    def _summary_lines(
        self,
        *,
        session_id: str,
        platform_family: str,
        metrics: dict[str, object],
        filter_summary: dict[str, object],
        narrative_lines: list[str],
        pareto_quick_analysis_rows: list[dict[str, object]],
    ) -> list[str]:
        platform_value = filter_summary.get("platform_family", [])
        if isinstance(platform_value, list):
            filtered_platforms = ", ".join(platform_value) or platform_family
        else:
            filtered_platforms = str(platform_value or platform_family)

        lines = [
            f"Session ID: {session_id}",
            f"Platform scope: {platform_family}",
            f"Filtered platforms: {filtered_platforms}",
            f"Devices in scope: {metrics.get('devices_in_scope', 0)}",
            f"Impacted devices: {metrics.get('impacted_devices_in_scope', 0)}",
            f"Bug findings: {metrics.get('bug_findings_in_scope', 0)}",
            f"Affected findings: {metrics.get('affected_findings_in_scope', 0)}",
            f"Fixed in target: {metrics.get('fixed_in_target_in_scope', 0)}",
            f"Top risk device: {metrics.get('top_risk_device', '-') or '-'}",
            "",
            "Active criteria:",
            *format_filter_lines(filter_summary),
            "",
            "Narrative:",
        ]
        if narrative_lines:
            lines.extend([f"- {line}" for line in narrative_lines[:5]])
        else:
            lines.append("- No prioritized narrative is available for the current filters.")
        lines.extend(["", "Pareto quick analysis:"])
        if pareto_quick_analysis_rows:
            for row in pareto_quick_analysis_rows[:3]:
                lines.append(
                    f"- {row.get('bug_id', '')}: risk {row.get('total_risk_score', 0)}, "
                    f"findings {row.get('finding_count', 0)}, remediation {row.get('remediation_status', '-')}"
                )
        else:
            lines.append("- No pareto quick-analysis rows are available for the current filters.")
        return lines

    def _chart_summary_lines(self, filter_summary: dict[str, object]) -> list[str]:
        return [
            "This chart reflects the active dashboard filters.",
            *format_filter_lines(filter_summary),
        ]

    def _add_textbox(
        self,
        *,
        slide: object,
        x: Emu,
        y: Emu,
        cx: Emu,
        cy: Emu,
        lines: list[str],
        font_size: Pt,
        bold: bool,
    ) -> None:
        textbox = slide.shapes.add_textbox(x, y, cx, cy)
        text_frame = textbox.text_frame
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        text_frame.clear()

        safe_lines = lines or [""]
        for index, line in enumerate(safe_lines):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = str(line)
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = font_size
                run.font.bold = bold
                run.font.name = "Arial"

